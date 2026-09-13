"""E5 — паттерны формул по блокам и месяцам (v2: абсолютные колонки внутри блока → C$[k],
обёртка IF($M>LAST_CLOSED_DATE,"",…) снимается для сравнения ядра)."""
import openpyxl, re, json, collections, datetime
from openpyxl.utils import get_column_letter as L, column_index_from_string as CI
from xl import real_formula

wb = openpyxl.load_workbook('book_2026-09-13.xlsx')
ws = wb['WB_Юнит_2025']
sec = json.load(open('audit_main.json'))['sections']

REF = re.compile(r"(?<![A-Za-z_\d!'])(\$?)([A-Z]{1,3})(\$?)(\d{1,7})(?![\d(])")
LCD = re.compile(r'^=IF\(R\[0\]C\$?\[0\]>LAST_CLOSED_DATE,"",(.*)\)$', re.S)

def norm(formula, r, bc, blockwidth=24):
    def rep(m):
        cabs, col, rabs, row = m.groups()
        ci = CI(col)
        dc = ci - bc
        if 0 <= dc < blockwidth or not cabs:
            cs = f"C{'$' if cabs else ''}[{dc}]"
        else:
            cs = f"C${col}"
        rs = f"R${row}" if rabs else f"R[{int(row)-r}]"
        return rs + cs
    return REF.sub(rep, formula)

def core(p):
    m = LCD.match(p)
    return m.group(1) if m else p

HDR = {off: ws.cell(736, 13 + off).value for off in range(24)}
result = {}
for s in sec:
    per_block = {}
    for (bc, title) in s['blocks']:
        offs = {}
        for off in range(24):
            pats = collections.Counter(); vals = 0; empt = 0
            for r in range(s['first'], s['last'] + 1):
                v = ws.cell(r, bc + off).value
                f, _ = real_formula(v)
                if f is not None: pats[core(norm(f, r, bc))] += 1
                elif v is None or v == '': empt += 1
                else: vals += 1
            offs[off] = dict(pats=pats.most_common(), vals=vals, empt=empt)
        per_block[title] = offs
    result[s['title']] = per_block

json.dump(result, open('audit_patterns.json', 'w'), ensure_ascii=False, indent=0, default=str)

# --- отчёт 1: сентябрь — согласованность блоков по каждому смещению
print('=== СЕНТЯБРЬ 2026: ядро формулы по смещению — сколько разных вариантов среди 24 блоков ===')
sep = result['Сентябрь 2026']
sep_core = {}
for off in range(24):
    variants = collections.Counter()
    for title, offs in sep.items():
        for p, n in offs[off]['pats']: variants[p] += n
    sep_core[off] = variants
    if variants:
        top = variants.most_common()
        print(f"off {off:2} {str(HDR[off])[:26]:26} variants={len(top)} | " + ' || '.join(f"{n}× {p[:110]}" for p, n in top[:3]))

# --- отчёт 2: все месяцы — доля ячеек, совпадающих с доминирующим сентябрьским ядром
print('\n=== ВСЕ МЕСЯЦЫ: доля формул блока, совпадающих с сентябрьским ядром (по смещению) ===')
hdrline = 'месяц           ' + ' '.join(f"{o:>4}" for o in range(24))
print(hdrline)
for s in sec:
    cells = []
    for off in range(24):
        ref = sep_core[off].most_common(1)[0][0] if sep_core[off] else None
        tot = 0; same = 0; vals = 0
        for title, offs in result[s['title']].items():
            for p, n in offs[off]['pats']:
                tot += n; same += n if p == ref else 0
            vals += offs[off]['vals']
        if ref is None: cells.append('  · ' if tot == 0 else f"{tot:>3}f")
        elif tot == 0: cells.append(' val' if vals else '  - ')
        else: cells.append(f"{100*same//tot:>3}%")
    print(f"{s['title']:16}" + ' '.join(cells))

# --- отчёт 3: месяцы — какие ядра встречаются на ключевых смещениях (7,10,12,13,20,22)
print('\n=== ВАРИАНТЫ ЯДРА ПО МЕСЯЦАМ (смещения 7 остатки, 10 доходность, 12 блогеры, 13 ДРР, 20 хранение, 22 доходность 1 шт) ===')
for off in (7, 10, 12, 13, 20, 22):
    print(f"\n-- off {off} {HDR[off]}")
    seen = collections.OrderedDict()
    for s in sec:
        variants = collections.Counter()
        for title, offs in result[s['title']].items():
            for p, n in offs[off]['pats']: variants[p] += n
        for p, n in variants.most_common():
            seen.setdefault(p, []).append(f"{s['title'][:3]}{s['title'][-2:]}:{n}")
    for i, (p, where) in enumerate(seen.items()):
        print(f"  V{i+1} [{sum(int(w.split(':')[1]) for w in where):5}] {p[:150]}")
        print(f"       {' '.join(where)}")
