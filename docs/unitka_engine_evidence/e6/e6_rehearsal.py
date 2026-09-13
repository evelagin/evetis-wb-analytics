"""E6-A — офлайн-репетиция на экспорте книги (13.09.2026 10:56 UTC). Ничего не пишет в книгу.
Выход: e6_dryrun_*.json + e6_dryrun_report.md — ожидаемый diff и данные для отката/QA."""
import openpyxl, re, json, collections, hashlib, datetime, sys
from openpyxl.utils import range_boundaries, get_column_letter as L, column_index_from_string as CI
sys.path.insert(0, '/mnt/user-data/outputs/e5')
from xl import real_formula

BOOK = '/mnt/user-data/outputs/e5/book_2026-09-13.xlsx'
wbf = openpyxl.load_workbook(BOOK)
wbv = openpyxl.load_workbook(BOOK, data_only=True)
ws = wbf['WB_Юнит_2025']
sec = json.load(open('/mnt/user-data/outputs/e5/audit_main.json'))['sections']
SEP = [s for s in sec if s['title'] == 'Сентябрь 2026'][0]
HIST = [s for s in sec if s['title'] != 'Сентябрь 2026']
out = {'generated_at': datetime.datetime.utcnow().isoformat() + 'Z', 'source_export': 'book_2026-09-13.xlsx (10:56 UTC)'}

def section_of(r):
    if r < 47: return 'top'
    for s in sec:
        if s['top'] <= r <= s['end']: return s['title']
    return '?'

# ---------------- A. Условное форматирование, главный лист ----------------
rules = []
for rng in ws.conditional_formatting:
    parts = str(rng.sqref).split()
    for rule in rng.rules:
        b = [range_boundaries(p) for p in parts]
        r1 = min(x[1] for x in b); r2 = max(x[3] for x in b)
        if rule.type == 'colorScale':
            cs = rule.colorScale
            sig = ('gradient', tuple((v.type, v.val) for v in cs.cfvo), tuple(c.rgb for c in cs.color))
        else:
            f = str(rule.formula[0]) if rule.formula else ''
            d = rule.dxf
            fill = d.fill.fgColor.rgb if d and d.fill and d.fill.fgColor is not None else None
            font = d.font.color.rgb if d and d.font and d.font.color is not None else None
            sig = ('boolean', f, fill, font)
        rules.append(dict(priority=rule.priority, parts=parts, bounds=b, r1=r1, r2=r2, sig=sig,
                          section=section_of(r1), type=rule.type, cells=sum((x[2]-x[0]+1)*(x[3]-x[1]+1) for x in b)))
rules.sort(key=lambda x: x['priority'])
hist_rules = [r for r in rules if r['r2'] <= SEP['top'] - 1]
# правила, у которых хотя бы одна часть в сентябре, остаются целиком (в т. ч. 6 градиентов с частями
# в августе: их домен = август + сентябрь, резать нельзя — семантика 1:1)
sep_rules = [r for r in rules if r['r2'] >= SEP['top']]
cross = [r for r in sep_rules if r['r1'] < SEP['top']]

# A1. сентябрь: точные дубликаты градиентов и полностью затенённые
seen = {}
keep, dup = [], []
covered = collections.defaultdict(set)   # col -> set of rows covered by earlier gradient rules
for r in sep_rules:
    if r['type'] != 'colorScale':
        keep.append(r); continue
    key = (tuple(sorted(r['parts'])), r['sig'])
    if key in seen:
        dup.append(dict(priority=r['priority'], parts=r['parts'], same_as=seen[key])); continue
    cells = set()
    for (c1, rr1, c2, rr2) in r['bounds']:
        for c in range(c1, c2 + 1):
            for rr in range(rr1, rr2 + 1): cells.add((c, rr))
    if all(rr in covered[c] for (c, rr) in cells):
        dup.append(dict(priority=r['priority'], parts=r['parts'], same_as='shadowed')); continue
    for (c, rr) in cells: covered[c].add(rr)
    seen[key] = r['priority']; keep.append(r)
sep_gradient_keep = [r for r in keep if r['type'] == 'colorScale']
sep_bool = [r for r in keep if r['type'] != 'colorScale']

# A2. формульные правила сентября: шаблоны и возможность объединения
ABS_IN_BLOCK = re.compile(r'\$([A-Z]{1,3})(\$?)(\d+)')
def relativize(formula):
    """$UX737 -> UX737 ; $UX$737 -> UX$737 ; глобальные ($WB$736 — зеркало LCD) не трогаем."""
    def rep(m):
        col = m.group(1)
        if col == 'WB': return m.group(0)
        return f"{col}{m.group(2)}{m.group(3)}"
    return ABS_IN_BLOCK.sub(rep, formula)
def template(sig):
    f = sig[1]
    t = re.sub(r'\$?[A-Z]{1,3}\$?\d+', 'REF', f); t = re.sub(r'\d+(\.\d+)?', '#', t)
    return (t, sig[2], sig[3])
groups = collections.OrderedDict()
for r in sep_bool:
    groups.setdefault(template(r['sig']), []).append(r)
bool_plan = []
for t, rs in groups.items():
    single_col = all(c1 == c2 for r in rs for (c1, _, c2, _) in r['bounds'])
    same_rows = len({(rr1, rr2) for r in rs for (_, rr1, _, rr2) in r['bounds']}) == 1
    has_global = '$WB$' in rs[0]['sig'][1]
    abs_refs = [m.group(0) for m in ABS_IN_BLOCK.finditer(rs[0]['sig'][1]) if m.group(1) != 'WB']
    if len(rs) == 1:
        action = 'keep'                       # уже одно правило (мульти-диапазон R7.5)
    elif single_col and same_rows:
        action = 'merge-relative' if abs_refs else 'merge'
    else:
        action = 'keep-per-block'
    example = rs[0]['sig'][1]
    bool_plan.append(dict(template=t[0], fill=t[1], font=t[2], rules_now=len(rs), parts_now=sum(len(r['parts']) for r in rs),
                          action=action, rules_after=1 if action.startswith('merge') else len(rs),
                          formula_example=example, formula_after=relativize(example) if action == 'merge-relative' else example,
                          ranges_after=[p for r in rs for p in r['parts']] if action.startswith('merge') else None))
sep_after_rules = len(sep_gradient_keep) + sum(b['rules_after'] for b in bool_plan)
sep_after_parts = sum(len(r['parts']) for r in sep_gradient_keep) + sum((len(b['ranges_after']) if b['ranges_after'] else b['parts_now']) for b in bool_plan)

# A3. история: печь фон, удалить правила
hist_by_sec = collections.Counter(r['section'] for r in hist_rules)
hist_parts = sum(len(r['parts']) for r in hist_rules)
bake = []
for s in [dict(title='top', top=1, end=46)] + HIST:
    bake.append(dict(section=s['title'], rows=[s['top'], s['end']], cols=[1, 600],
                     rules=hist_by_sec.get(s['title'], 0), cells=(s['end'] - s['top'] + 1) * 600))
out['cf_main'] = dict(
    now=dict(rules=len(rules), parts=sum(len(r['parts']) for r in rules), gradient=sum(1 for r in rules if r['type'] == 'colorScale'),
             boolean=sum(1 for r in rules if r['type'] != 'colorScale')),
    history=dict(rules=len(hist_rules), parts=hist_parts, bake_sections=bake),
    september=dict(rules_now=len(sep_rules), parts_now=sum(len(r['parts']) for r in sep_rules),
                   gradient_now=sum(1 for r in sep_rules if r['type'] == 'colorScale'), gradient_dropped=len(dup),
                   gradient_after=len(sep_gradient_keep), boolean_now=len(sep_bool), boolean_plan=bool_plan,
                   rules_after=sep_after_rules, parts_after=sep_after_parts),
    after=dict(rules=sep_after_rules, parts=sep_after_parts),
    dropped_gradient_examples=dup[:10],
    cross_boundary_rules_kept=[dict(priority=r['priority'], parts=r['parts']) for r in cross],
)

# другие листы
cf_other = {}
for w in wbf.worksheets:
    if w.title == 'WB_Юнит_2025': continue
    n = sum(len(rng.rules) for rng in w.conditional_formatting)
    parts = sum(len(str(rng.sqref).split()) * len(rng.rules) for rng in w.conditional_formatting)
    cf_other[w.title] = dict(rules=n, parts=parts, action=('keep' if w.title == 'Склад' else ('bake+delete' if n else 'none')))
out['cf_other'] = cf_other

# ---------------- B. Правки F1–F6 ----------------
wv = wbv['WB_Юнит_2025']; w24 = wbf['WB_Юнит_2024']; v24 = wbv['WB_Юнит_2024']
def cur(sheet_f, sheet_v, a1):
    f, _ = real_formula(sheet_f[a1].value)
    return dict(formula=f, value=sheet_v[a1].value if not isinstance(sheet_v[a1].value, datetime.datetime) else str(sheet_v[a1].value))
fixes = []
for a1 in ('C3', 'BU61', 'BU96', 'BU130'):
    fixes.append(dict(id='F1', sheet='WB_Юнит_2024', cell=a1, before=cur(w24, v24, a1), after='без изменений формулы (лист «Блогеры» восстановлен из резервной копии 1H8ldc…, values-only)', dependents=['C4', 'BU10', 'BV10', 'BU62', 'BU97', 'BU131']))
fixes.append(dict(id='F1', sheet='WB_Юнит_2025', cell='VW10', before=cur(ws, wv, 'VW10'), after='без изменений формулы (лист «Блогеры» восстановлен)', dependents=['VW44', 'E3', 'E4']))
fixes.append(dict(id='F2', sheet='WB_Юнит_2025', cell='DP293', before=cur(ws, wv, 'DP293'), after='пусто', issue='ADS_MISSING 2025-08-05 305101361', dependents=['DO293', 'I293', 'J293', 'I320', 'J320', 'DO320', 'VR289', 'VR290', 'VR46', 'VS46', 'DN44', 'DN45']))
fixes.append(dict(id='F3', sheet='WB_Юнит_2025', cell='BD663', before=cur(ws, wv, 'BD663'), after='=BD633*AO663'))
fixes.append(dict(id='F3', sheet='WB_Юнит_2025', cell='CB663', before=cur(ws, wv, 'CB663'), after='=CB633*BM663'))
fixes.append(dict(id='F4', sheet='WB_Юнит_2025', cell='JX760', before=cur(ws, wv, 'JX760'), after='=IF($JR760>LAST_CLOSED_DATE,"",JX759-JU760+JW759)' if 'LAST_CLOSED_DATE' in (real_formula(ws['JX760'].value)[0] or '') else '=JX759-JU760+JW759'))
june = [s for s in sec if s['title'] == 'Июнь 2026'][0]
f5 = []
for c in range(3, 11):
    a1 = f"{L(c)}{june['first']}"
    f, _ = real_formula(ws[a1].value)
    f5.append(dict(cell=f"{L(c)}{june['first']}:{L(c)}{june['last']}", before=f, after=f.replace(':SX', ':UT').replace(':SY', ':UU').replace(':SZ', ':UV').replace(':TA', ':UW').replace(':TB', ':UX').replace(':TC', ':UY').replace(':TG', ':VC').replace(':TH', ':VD') if f else None))
fixes.append(dict(id='F5', sheet='WB_Юнит_2025', cells=f5, note='диапазон FILTER до 24-го блока (UT..VD), как в июле–сентябре; значения не меняются (у блоков 23–24 в июне заказов 0)'))
fixes.append(dict(id='F6', sheet='WB_Юнит_2025', cell='TD765', before=cur(ws, wv, 'TD765'), after='без изменений (D4: сохранить как предполагаемую ручную поправку; ISSUE UNEXPLAINED_MANUAL_ADJUSTMENT до сверки)'))
fixes.append(dict(id='F6', sheet='WB_Юнит_2025', cell='TD730', before=cur(ws, wv, 'TD730'), after='без изменений (D4)'))
out['fixes'] = fixes

# ---------------- C. Заморозка ----------------
def sheet_stats(name):
    wf = wbf[name]; wvv = wbv[name]
    forms = 0; errs = 0; today = 0
    for row in wf.iter_rows():
        for c in row:
            f, _ = real_formula(c.value)
            if f:
                forms += 1
                if 'TODAY(' in f: today += 1
                cv = wvv[c.coordinate].value
                if isinstance(cv, str) and cv.startswith('#'): errs += 1
    return dict(formulas=forms, error_cells=errs, today=today, max_row=wf.max_row, max_col=wf.max_column)
out['freeze'] = {'Акции': dict(**sheet_stats('Акции'), action='formulas→values, A1 TODAY→дата, A2 пометка АРХИВ, hide'),
                 'OZON_Юнит_2025': dict(**sheet_stats('OZON_Юнит_2025'), action='formulas→values, ошибки→пусто, УФ bake+delete, hide'),
                 'WB_Юнит_2024': dict(**sheet_stats('WB_Юнит_2024'), action='только УФ bake+delete (формулы остаются: витрина E2/E3 2025 ссылается на C2/C3)')}

# ---------------- D. Сетка ----------------
trim = {}
for w in wbf.worksheets:
    last_r = 0; last_c = 0
    for row in w.iter_rows():
        for c in row:
            if c.value not in (None, ''):
                last_r = max(last_r, c.row); last_c = max(last_c, c.column)
    trim[w.title] = dict(max_row=w.max_row, max_col=w.max_column, last_used_row=last_r, last_used_col=last_c,
                         delete_cols=(w.max_column - last_c) if w.max_column - last_c > 0 else 0,
                         delete_rows=max(0, w.max_row - (last_r + 50)))
trim['WB_Юнит_2025']['note'] = 'используются колонки до WB (600): сводки VR:VW и зеркало LCD; удалить WC:XZ'
out['trim'] = trim

# ---------------- E. Базовые отпечатки для QA ----------------
def hash_sheet(name, formulas=False):
    w = wbf[name] if formulas else wbv[name]
    h = hashlib.sha256(); n = 0
    for row in w.iter_rows():
        for c in row:
            v = c.value
            if v is None: continue
            if formulas:
                f, _ = real_formula(v)
                if f is None: continue
                v = f
            n += 1
            h.update(f"{c.coordinate}={v!r};".encode())
    return dict(cells=n, sha256=h.hexdigest()[:16])
out['baseline'] = {w.title: dict(values=hash_sheet(w.title), formulas=hash_sheet(w.title, True)) for w in wbf.worksheets}

# ---------------- F. Блогеры ----------------
out['bloggers_restore'] = dict(source='1H8ldcUNpMbkceVzfHGTNwdMLqdaNVzP8JJ1qK_i3Z6s («…_BACKUP_до_аудита», 10.09.2026 14:58 UTC)',
                               evidence='ZZ_AUDIT_1789053792496.json (снимок 10.09 15:23 UTC): лист «Блогеры» i=9, 2451×21, формул 3018, непустых 8642, ошибок 0',
                               also_removed_sheets=['Поставки FBO', 'Оборачиваемость товара', 'Финальная сс', 'реклама ОЗОН', 'реклама (динамика)', 'тесты CTR'],
                               method='values-only копия в лист «Блогеры» живой книги; формулы 2024!C3/BU61/BU96/BU130 и 2025!VW10 не меняются; провенанс RESTORED_FROM_BACKUP в A1-заметке и ZZ_AUDIT_LOG',
                               referenced_cells=['I3', 'I4', 'I5', 'K13'])

json.dump(out, open('e6_dryrun.json', 'w'), ensure_ascii=False, indent=1, default=str)

# ---------------- отчёт ----------------
m = out['cf_main']
lines = [f"# E6-A dry-run (офлайн, экспорт {out['source_export']})", '',
         '## УФ главного листа',
         f"сейчас: {m['now']['rules']} правил / {m['now']['parts']} фрагментов (градиентов {m['now']['gradient']}, формульных {m['now']['boolean']})",
         f"история + строки 1–46: {m['history']['rules']} правил / {m['history']['parts']} фрагментов → 0 (фон печётся статически, {len(m['history']['bake_sections'])} секций)",
         f"сентябрь: {m['september']['rules_now']} правил / {m['september']['parts_now']} фрагментов → {m['september']['rules_after']} правил / {m['september']['parts_after']} фрагментов "
         f"(градиенты {m['september']['gradient_now']} → {m['september']['gradient_after']}: сброшено {m['september']['gradient_dropped']} дублей/затенённых; формульные {m['september']['boolean_now']} → {sum(b['rules_after'] for b in m['september']['boolean_plan'])})",
         f"**итого главный лист: {m['now']['rules']} → {m['after']['rules']} правил, {m['now']['parts']} → {m['after']['parts']} фрагментов**", '',
         '| шаблон формульного правила | сейчас правил/фрагментов | действие | после |', '|---|---|---|---|']
for b in m['september']['boolean_plan']:
    lines.append(f"| `{b['template'][:70]}` | {b['rules_now']}/{b['parts_now']} | {b['action']} | {b['rules_after']} |")
lines += ['', '## УФ других листов'] + [f"- {k}: {v['rules']} правил / {v['parts']} фрагментов → {v['action']}" for k, v in out['cf_other'].items() if v['rules']]
lines += ['', '## Правки F1–F6 (до/после)']
for f in fixes:
    if 'cells' in f:
        lines.append(f"- **{f['id']}** {f['sheet']}: {len(f['cells'])} колонок сводки июня — {f['note']}")
        for c in f['cells']: lines.append(f"    - `{c['cell']}`: `{(c['before'] or '')[:80]}` → `{(c['after'] or '')[:80]}`")
    else:
        lines.append(f"- **{f['id']}** {f['sheet']}!{f['cell']}: `{str(f['before']['formula'] or f['before']['value'])[:90]}` (значение `{f['before']['value']}`) → {f['after']}" + (f" · ISSUE `{f['issue']}`" if f.get('issue') else ''))
lines += ['', '## Заморозка'] + [f"- {k}: формул {v['formulas']}, ошибок {v['error_cells']}, TODAY {v['today']} → {v['action']}" for k, v in out['freeze'].items()]
lines += ['', '## Сетка'] + [f"- {k}: {v['max_row']}×{v['max_col']}, использовано до {v['last_used_row']}×{v['last_used_col']} → удалить колонок {v['delete_cols']}, строк {v['delete_rows']}" for k, v in trim.items()]
lines += ['', '## Блогеры (D2)', f"- источник: {out['bloggers_restore']['source']}", f"- доказательство: {out['bloggers_restore']['evidence']}", f"- метод: {out['bloggers_restore']['method']}"]
lines += ['', '## Базовые отпечатки (значения/формулы по листам, для сверки с живым снимком)'] + [f"- {k}: values {v['values']['cells']} ячеек `{v['values']['sha256']}`, formulas {v['formulas']['cells']} `{v['formulas']['sha256']}`" for k, v in out['baseline'].items()]
open('e6_dryrun_report.md', 'w').write('\n'.join(lines) + '\n')
print('\n'.join(lines))
