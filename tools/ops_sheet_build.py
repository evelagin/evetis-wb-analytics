#!/usr/bin/env python3
"""EVETIS OPERATIONS (Stage C1) — сборка книги для Google Sheets из представлений evetis_ops.

Лист только показывает значения: вся логика — в BigQuery (sql/evetis_ops/ops_08_supply_plan.sql).
В книге нет формул, нет кнопок, нет записи в BigQuery.

Запуск:  python3 tools/ops_sheet_build.py [<путь к .xlsx>]
Нужен openpyxl:  python3 -m venv .venv && .venv/bin/pip install openpyxl
Загрузка в Google Sheets: Файл → Импорт → Загрузить → «Заменить таблицу» (URL книги не меняется).
"""
import sys
import os
from datetime import datetime, timezone, timedelta

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import ops_deploy as d  # noqa: E402

from openpyxl import Workbook  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

P = d.P
FONT = 'Arial'
C_DARK, C_WHITE, C_MUTED, C_LINE = '434343', 'FFFFFF', 'B7B7B7', 'D9D9D9'
C_WB, C_OZ, C_BUNDLE, C_PICK = '4A148C', '0D47A1', '7F6000', '274E13'
ST = {   # статус → (заливка, цвет текста, жирный)
    'SHIP': ('E6F4EA', '274E13', True), 'WAIT': ('FFF2CC', '7F6000', False),
    'REVIEW': ('FCE8E6', 'A61C00', True), 'FF LIMIT': ('FCE8E6', 'A61C00', False),
    'OK': ('EFEFEF', '434343', False), 'NO PLAN': ('FFFFFF', 'B7B7B7', False),
}
INT, F1, DATE = '#,##0', '#,##0.0', 'dd.mm.yyyy'
THIN = Side(style='thin', color=C_LINE)


def query(sql):
    """Строки представления как список словарей с типами BigQuery."""
    ok, msg, _ = True, '', None
    r = d.call('POST', f'{d.API}/jobs', {
        'configuration': {'query': {'query': sql, 'useLegacySql': False},
                          'labels': {'stage': 'stage-c1', 'sandbox': 'false'}},
        'jobReference': {'projectId': P, 'location': 'EU'}})
    if 'error' in r:
        raise SystemExit('BQ insert: ' + str(r['error'])[:400])
    jid = r['jobReference']['jobId']
    import time
    while True:
        j = d.call('GET', f'{d.API}/jobs/{jid}?location=EU')
        if j.get('status', {}).get('state') == 'DONE':
            break
        time.sleep(2)
    if j['status'].get('errorResult'):
        raise SystemExit('BQ query: ' + j['status']['errorResult'].get('message', '')[:500])
    rows, token = [], None
    while True:
        url = f'{d.API}/queries/{jid}?location=EU&maxResults=2000' + (f'&pageToken={token}' if token else '')
        res = d.call('GET', url)
        fields = res.get('schema', {}).get('fields', [])
        for row in res.get('rows', []):
            rec = {}
            for f, cell in zip(fields, row['f']):
                v = cell.get('v')
                t = f['type']
                if v is None:
                    rec[f['name']] = None
                elif t in ('INTEGER', 'INT64'):
                    rec[f['name']] = int(v)
                elif t in ('FLOAT', 'FLOAT64', 'NUMERIC', 'BIGNUMERIC'):
                    rec[f['name']] = float(v)
                elif t in ('BOOLEAN', 'BOOL'):
                    rec[f['name']] = (v == 'true')
                elif t == 'TIMESTAMP':
                    rec[f['name']] = datetime.fromtimestamp(float(v), tz=timezone.utc).astimezone(
                        timezone(timedelta(hours=3))).strftime('%Y-%m-%d %H:%M')
                else:
                    rec[f['name']] = v
            rows.append(rec)
        token = res.get('pageToken')
        if not token:
            return rows


class Sheet:
    """Тонкая обёртка над листом: заголовки, полосы разделов, строки таблиц."""

    def __init__(self, ws):
        self.ws = ws
        self.r = 1

    def title(self, text, sub=None, width=21):
        ws = self.ws
        ws.cell(self.r, 1, text).font = Font(FONT, 15, bold=True, color=C_DARK)
        ws.merge_cells(start_row=self.r, start_column=1, end_row=self.r, end_column=width)
        ws.row_dimensions[self.r].height = 24
        self.r += 1
        for line in (sub or []):
            c = ws.cell(self.r, 1, line)
            c.font = Font(FONT, 9, color='666666')
            ws.merge_cells(start_row=self.r, start_column=1, end_row=self.r, end_column=width)
            self.r += 1
        self.r += 1

    def band(self, text, color, width=21, note=None, note_as_comment=False):
        ws = self.ws
        c = ws.cell(self.r, 1, text)
        c.font = Font(FONT, 11, bold=True, color=C_WHITE)
        c.fill = PatternFill('solid', fgColor=color)
        c.alignment = Alignment(vertical='center', indent=1)
        if note and note_as_comment:
            c.comment = Comment(note, 'evetis_ops')
            note = None
        for col in range(2, width + 1):
            ws.cell(self.r, col).fill = PatternFill('solid', fgColor=color)
        ws.merge_cells(start_row=self.r, start_column=1, end_row=self.r, end_column=width)
        ws.row_dimensions[self.r].height = 22
        self.r += 1
        if note:
            c = ws.cell(self.r, 1, note)
            c.font = Font(FONT, 9, italic=True, color='666666')
            ws.merge_cells(start_row=self.r, start_column=1, end_row=self.r, end_column=width)
            self.r += 1
        return self.r - 1

    def header(self, cols):
        ws = self.ws
        for i, col in enumerate(cols, start=1):
            c = ws.cell(self.r, i, col['h'])
            c.font = Font(FONT, 9, bold=True, color=C_WHITE)
            c.fill = PatternFill('solid', fgColor=C_DARK)
            c.alignment = Alignment(wrap_text=True, vertical='center',
                                    horizontal='left' if i <= 2 else 'center')
            c.border = Border(bottom=THIN)
            if col.get('note'):
                c.comment = Comment(col['note'], 'evetis_ops')
            ws.column_dimensions[get_column_letter(i)].width = col['w']
        ws.row_dimensions[self.r].height = 34
        self.r += 1

    def rows(self, cols, data, status_key=None, bold_key=None, total=None, color_key='status_code'):
        ws = self.ws
        for rec in data:
            code = rec.get(color_key) if status_key else None
            dim = code == 'NO PLAN'
            for i, col in enumerate(cols, start=1):
                v = rec.get(col['k'])
                c = ws.cell(self.r, i, v)
                c.font = Font(FONT, 10, color=C_MUTED if dim else C_DARK)
                c.border = Border(bottom=THIN)
                if col.get('f'):
                    c.number_format = col['f']
                    c.alignment = Alignment(horizontal='right')
                if i <= 2 or col['k'] in ('reason', 'status_label', 'bom_text', 'stage'):
                    c.alignment = Alignment(horizontal='left', vertical='center')
                if status_key and col['k'] == status_key:
                    fill, fg, bold = ST.get(code or '', ('FFFFFF', C_DARK, False))
                    c.fill = PatternFill('solid', fgColor=fill)
                    c.font = Font(FONT, 9, bold=bold, color=fg)
                    c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=False)
                if bold_key and col['k'] == bold_key:
                    on = bool(v)
                    c.font = Font(FONT, 11, bold=on, color='0B5394' if on else C_MUTED)
                    c.fill = PatternFill('solid', fgColor='E8F0FE' if on else 'FFFFFF')
            ws.row_dimensions[self.r].height = 19
            self.r += 1
        if total:
            for i, col in enumerate(cols, start=1):
                v = total.get(col['k'])
                c = ws.cell(self.r, i, v)
                c.font = Font(FONT, 10, bold=True, color=C_DARK)
                c.fill = PatternFill('solid', fgColor='F3F3F3')
                c.border = Border(top=Side(style='medium', color=C_DARK))
                if col.get('f'):
                    c.number_format = col['f']
            self.r += 1
        self.r += 1


def col(h, k, w, f=None, note=None):
    return {'h': h, 'k': k, 'w': w, 'f': f, 'note': note}


PLAN_COLS = [
    col('Товар', 'product_name', 32),
    col('SKU', 'card_sku', 26),
    col('На площадке\nсейчас', 'on_marketplace', 11, INT,
        'ON_WB / ON_OZON — остаток на складе площадки по API. Резерв на ФФ сюда НЕ входит.'),
    col('В пути', 'in_transit', 9, INT, 'Поставка уехала и видна площадке как «в пути».'),
    col('На приёмке\nOzon', 'in_acceptance', 10, INT,
        'IN_ACCEPTANCE — поставка на приёмке Ozon. Принятое количество Ozon ещё не сообщил. '
        'К остатку площадки не прибавляется; считается входящим ровно один раз.'),
    col('Резерв\nна ФФ', 'reserved_on_ff', 10, INT,
        'RESERVED_OZON_ON_FF / резерв WB — физически лежит на ФФ и уже обещано этой площадке.'),
    col('Прочее\nвходящее', 'other_inbound', 10, INT,
        'Резерв в кабинете площадки без отметки ФФ + уехало, но площадка ещё не видит.'),
    col('Итого\nвходящее', 'committed_inbound', 10, INT,
        'Сумма всего входящего. Ни одна единица не посчитана дважды (класс сверки заказа — один).'),
    col('Спрос до\nприбытия', 'demand_until_arrival', 10, INT,
        'План продаж с сегодня до даты прибытия: сборка на ФФ (2 рабочих дня) + срок канала.'),
    col('Прогноз\nк прибытию', 'projected_at_arrival', 11, INT,
        '= на площадке + входящее − спрос до прибытия. Если ушло бы в минус, показан 0, '
        'а дефицит назван в колонке «Почему».'),
    col('План/день', 'daily_plan', 9, F1, 'Средний план продаж в день в окне покрытия после прибытия.'),
    col('Факт/день\n30 дн', 'actual_daily_30d', 9, F1,
        'Фактические заказы в день за последние 30 дней. Только для сравнения, в расчёте не участвует.'),
    col('Цель\nк прибытию', 'target_at_arrival', 10, INT,
        '= план/день × (целевое покрытие + страховые дни). 45 дней считаются ПОСЛЕ прибытия.'),
    col('в т.ч.\nстраховой', 'safety_units', 9, INT, 'Страховой запас = план/день × страховые дни (WB 7, Ozon 10).'),
    col('Потребность', 'need_math', 11, INT, '= цель к прибытию − прогноз к прибытию, но не меньше 0.'),
    col('Кратность', 'multiple_applied', 9, INT,
        'Операционная кратность отгрузки (конвенция владельца, не заводской короб). Округление — вверх.'),
    col('План / Факт', 'plan_actual_ratio', 9, F1,
        'Во сколько раз план выше факта за 30 дней. Больше 2 — план агрессивный, объём требует подтверждения.'),
    col('Уверенность\nв спросе', 'demand_confidence', 11, None,
        'HIGH — план ≈ факт (до 1,25×); MEDIUM — до 2×; LOW — больше 2× или продаж за 30 дней не было.'),
    col('РЕКОМЕНДАЦИЯ,\nпозиций', 'rec_final', 13, INT,
        'Сколько отгрузить сейчас, в единицах продажи (позициях). 0 — ждать или сначала проверить.'),
    col('в физических\nединицах', 'rec_physical_units', 12, INT,
        'Та же рекомендация во флаконах / банках: позиции набора умножены на состав по BOM.'),
    col('Покрытие\nпосле, дн', 'resulting_cover_days', 10, F1,
        'На сколько дней хватит после прибытия. Выше порога (WB 66, Ozon 69) — статус ПРОВЕРИТЬ.'),
    col('ФФ свободно,\nпозиций', 'ff_free_cards_now', 11, INT,
        'Свободно на ФФ под эту позицию (для набора — комплектов). Общий пул WB и Ozon: '
        'строки обслуживаются по срочности, одна единица не обещана дважды.'),
    col('Статус', 'status_label', 30),
    col('Внимание', 'demand_warning', 44, None, 'Предупреждение владельцу. Рекомендацию не меняет.'),
    col('Почему', 'reason', 96),
]

TOP_COLS = [
    col('Товар', 'product_name', 34),
    col('Соло / Набор', 'kind', 12, None, 'Набор = одна позиция продажи из нескольких физических единиц.'),
    col('Позиции, шт.', 'rec_final', 12, INT, 'Проданные / отгружаемые позиции.'),
    col('Физические\nединицы, шт.', 'rec_physical_units', 13, INT, 'Флаконы / банки после раскладки по BOM.'),
    col('ФФ свободно,\nпозиций', 'ff_free_cards_now', 12, INT),
    col('Покрытие\nпосле, дн', 'resulting_cover_days', 11, F1),
    col('Статус', 'status_label', 30),
    col('Почему', 'reason_short', 110),
]

TASK_MOVE_COLS = [
    col('Товар', 'product_name', 34), col('SKU', 'internal_sku', 24),
    col('Снять под новые\nотгрузки, шт.', 'move_pallet_to_shelf_new', 15, INT),
    col('Снять под уже\nзарезервированное, шт.', 'move_pallet_to_shelf_reserved', 18, INT,
        'Резерв поставок Ozon от 07.09, который физически лежит на паллетах.'),
    col('ИТОГО снять\nс паллет, шт.', 'move_pallet_to_shelf_total', 15, INT),
    col('Полка после', 'shelf_after_pick', 11, INT), col('Паллеты после', 'pallet_after_pick', 12, INT),
]

TASK_BUILD_COLS = [
    col('Набор', 'product_name', 34), col('Состав', 'bom_text', 40),
    col('Собрать,\nнаборов', 'to_assemble_now', 11, INT),
    col('в т.ч. под резерв\nотгрузок', 'to_assemble_reserved', 15, INT),
    col('Физические\nединицы, шт.', 'assemble_physical_units', 13, INT),
    col('Лимитирующий\nкомпонент', 'limiting_after_plan', 20),
]

TASK_COMP_COLS = [
    col('Товар', 'product_name', 34), col('SKU', 'internal_sku', 24),
    col('Соло WB, шт.', 'solo_units_wb', 12, INT), col('Соло Ozon, шт.', 'solo_units_ozon', 12, INT),
    col('В наборы, шт.', 'bundle_component_units', 12, INT), col('FBS, шт.', 'fbs_component_units', 9, INT),
    col('ИТОГО физических\nединиц, шт.', 'total_physical_units', 16, INT),
    col('Уже зарезервировано\nна ФФ, шт.', 'reserved_units_already', 17, INT),
    col('Свободно на ФФ\nпосле операции, шт.', 'free_ff_after_operation', 18, INT),
]

BUNDLE_COLS = [
    col('Набор', 'product_name', 32), col('SKU', 'bundle_sku', 26),
    col('Состав', 'bom_text', 40),
    col('Мощность\nиз свободного', 'capacity_alone', 12, INT,
        'Сколько наборов можно собрать из свободного запаса, если другие наборы из тех же компонентов не собираются. '
        'Мощности разных наборов НЕ складываются.'),
    col('в т.ч.\nс полки', 'capacity_alone_shelf', 9, INT, 'Сколько можно собрать, ничего не снимая с паллет.'),
    col('План/день\nWB', 'plan_day_wb', 9, F1), col('План/день\nOzon', 'plan_day_ozon', 9, F1),
    col('К отгрузке\nWB', 'ship_wb', 10, INT, 'Рекомендация плана поставок для WB, позиций.'),
    col('К отгрузке\nOzon', 'ship_ozon', 10, INT, 'Рекомендация плана поставок для Ozon, позиций.'),
    col('Резерв\nFBS', 'fbs_reserve', 8, INT, 'FBS выключен (0) — поле сохранено для Stage E.'),
    col('Уже\nсобрано', 'assembled_sets', 9, INT, 'Собранные наборы на ФФ (FBO под отгрузки + FBS).'),
    col('Под резерв\nотгрузок', 'to_assemble_reserved', 11, INT,
        'Наборы в уже зарезервированных отгрузках (компоненты в резерве), которые ещё не собраны.'),
    col('СОБРАТЬ\nСЕЙЧАС', 'to_assemble_now', 12, INT,
        '= под резерв отгрузок + новые рекомендации WB и Ozon + резерв FBS − уже собрано.'),
    col('Лимитирующий\nкомпонент', 'limiting_after_plan', 20, None,
        'Компонент, которого после всего плана остаётся меньше всего.'),
    col('Мощность\nпосле плана', 'capacity_after_plan', 12, INT),
]

BOM_COLS = [
    col('Набор', 'product_name', 32), col('Компонент', 'component_sku', 24),
    col('На 1 набор', 'component_qty', 10, INT), col('Собрать, наборов', 'to_assemble_now', 12, INT),
    col('Физические\nединицы, шт.', 'units_total', 13, INT),
    col('из них уже\nв резерве', 'units_from_reserved', 12, INT,
        'Компоненты уже зарезервированы под существующие отгрузки — новый запас не нужен.'),
    col('из свободного\nзапаса', 'units_from_free', 13, INT,
        'Берётся из свободного запаса ФФ — эти единицы и попадают в «снятие с паллет».'),
]

PICK_COLS = [
    col('Компонент', 'product_name', 32), col('SKU', 'internal_sku', 24),
    col('Соло WB', 'solo_wb_need', 10, INT, 'Рекомендации плана по одиночным позициям WB.'),
    col('Соло Ozon', 'solo_ozon_need', 10, INT, 'Рекомендации плана по одиночным позициям Ozon.'),
    col('Компоненты\nнаборов', 'bundle_component_need', 12, INT, 'Раскладка рекомендованных наборов по BOM.'),
    col('FBS', 'fbs_component_need', 8, INT),
    col('ИТОГО\nспрос', 'total_physical_demand', 11, INT, 'Сумма — ровно то, что рекомендовано выше.'),
    col('Полка\nсвободно', 'shelf_free', 10, INT, 'Оперативное хранение за вычетом резерва Ozon из API без отметки ФФ.'),
    col('Паллеты\nсвободно', 'pallet_free', 11, INT),
    col('СНЯТЬ\nС ПАЛЛЕТ', 'to_pick_from_pallet', 12, INT, '= спрос − свободное на полке (в пределах паллет).'),
    col('Резерв на\nпаллетах', 'reserved_on_pallet_to_pick', 11, INT,
        'Уже зарезервировано под отгрузки Ozon от 07.09, но лежит на паллетах — тоже нужно снять.'),
    col('ИТОГО\nснять', 'pick_total_with_reserved', 11, INT),
    col('Полка\nпосле', 'shelf_after_pick', 9, INT, 'Не может быть отрицательной.'),
    col('Паллеты\nпосле', 'pallet_after_pick', 10, INT),
    col('Свободно\nпосле плана', 'free_after_plan', 12, INT),
]

FF_COLS = [
    col('Товар', 'product_name', 32), col('SKU', 'internal_sku', 24),
    col('PALLET FREE', 'pallet_free', 12, INT, 'Длительное хранение, свободно.'),
    col('SHELF FREE', 'shelf_free', 11, INT, 'Оперативное хранение (полка), свободно.'),
    col('RESERVED WB', 'reserved_wb', 12, INT),
    col('RESERVED OZON', 'reserved_ozon', 13, INT, 'Физически на ФФ под поставки Ozon. НЕ остаток площадки.'),
    col('в т.ч. на\nпаллетах', 'reserved_ozon_on_pallet', 10, INT, 'Эти единицы нужно снять с паллет перед отгрузкой.'),
    col('ASSEMBLED', 'assembled', 11, INT, 'Собранные наборы под конкретные отгрузки (FBO).'),
    col('FBS READY', 'fbs_ready', 10, INT),
    col('SHIPPED\nUNCONFIRMED', 'shipped_unconfirmed', 13, INT, 'Уехало с ФФ, площадка ещё не подтвердила. Вне итога ФФ.'),
    col('TOTAL PHYSICAL', 'total_physical', 13, INT, 'Физически на ФФ = паллеты + полка + резервы + собранное + FBS.'),
    col('IN_ACCEPTANCE\n(Ozon)', 'ozon_in_acceptance', 13, INT, 'Справочно: на приёмке Ozon, вне ФФ.'),
    col('ON_OZON', 'on_ozon', 10, INT, 'Справочно: остаток на складе Ozon (API).'),
    col('ON_WB', 'on_wb', 10, INT, 'Справочно: остаток на складе WB (API).'),
]

SHIP_COLS = [
    col('Канал', 'channel', 9), col('Номер', 'order_number', 16),
    col('Этап', 'stage', 36, None, 'Один заказ — один класс: резерв на ФФ, приёмка, в пути или принято.'),
    col('Статус API', 'api_state', 30), col('Статус\nжурнала', 'ledger_status', 11),
    col('Позиции, шт.', 'cards', 11, INT, 'Проданные / отгружаемые позиции (набор = одна позиция).'),
    col('в т.ч.\nнаборов', 'bundle_cards', 9, INT),
    col('Физические\nединицы, шт.', 'physical_units', 13, INT, 'После раскладки наборов по BOM.'),
    col('Резерв на ФФ,\nфиз. ед.', 'reserved_on_ff_units', 12, INT),
    col('с полки', 'reserved_shelf', 9, INT), col('с паллет', 'reserved_pallet', 9, INT),
    col('В пути /\nприёмка', 'pipeline_units', 10, INT), col('Принято', 'handed_over_units', 9, INT),
    col('Точка сдачи', 'dropoff', 18), col('План. дата', 'planned_date', 11, DATE),
    col('Создан', 'created_date', 11, DATE), col('Доказательство статуса на 09.09', 'evidence', 46),
]


def build(path):
    plan = query(f'SELECT * FROM `{P}.evetis_ops.V_OPS_SUPPLY_PLAN` ORDER BY sort_key, card_sku')
    for r in plan:
        r['other_inbound'] = (r['api_reserved_unconfirmed'] or 0) + (r['shipped_unconfirmed'] or 0)
    cal = query(f'SELECT * FROM `{P}.evetis_ops.V_OPS_SUPPLY_CALENDAR` ORDER BY channel DESC')
    bundles = query(f'SELECT * FROM `{P}.evetis_ops.V_OPS_BUNDLE_PRODUCTION` ORDER BY to_assemble_now DESC, bundle_sku')
    bom = query(f'SELECT * FROM `{P}.evetis_ops.V_OPS_BUNDLE_BOM_EXPANSION` ORDER BY bundle_sku, component_sku')
    pick = query(f'SELECT * FROM `{P}.evetis_ops.V_OPS_PICK_FROM_STORAGE` ORDER BY to_pick_from_pallet DESC, internal_sku')
    ff = query(f'SELECT * FROM `{P}.evetis_ops.V_OPS_FF_STOCK_SHEET` ORDER BY total_physical DESC, internal_sku')
    ships = query(f'SELECT * FROM `{P}.evetis_ops.V_OPS_SHIPMENTS_SHEET` ORDER BY stage, order_number')
    task = query(f'SELECT * FROM `{P}.evetis_ops.V_OPS_FF_TASK` ORDER BY move_pallet_to_shelf_total DESC, internal_sku')
    chan = query(f'SELECT * FROM `{P}.evetis_ops.REF_CHANNEL_SHIPPING` ORDER BY channel DESC')
    logi = query(f'''SELECT internal_sku, channel, shipment_multiple, shipment_multiple_low_demand, min_shipment_units,
                     factory_carton_qty, unit_weight_kg, unit_volume_l, box_weight_kg, fbs_reserve_units,
                     JSON_VALUE(field_provenance, '$.unit_weight_kg.class') AS weight_class,
                     JSON_VALUE(field_provenance, '$.shipment_multiple.class') AS multiple_class,
                     JSON_VALUE(field_provenance, '$.factory_carton_qty.hint') AS carton_hint,
                     effective_from, source
                     FROM `{P}.evetis_ops.REF_SKU_LOGISTICS` ORDER BY channel DESC, internal_sku''')
    cfg = query(f"SELECT config_key, config_value, note FROM `{P}.evetis_ops.OPS_CONFIG` ORDER BY config_key")
    fresh = query(f'''SELECT
      (SELECT FORMAT_DATE('%d.%m.%Y', MAX(_snapshot_date)) FROM `{P}.wb_raw.V_WB_STOCKS_T5_CURRENT`) AS wb_stock,
      (SELECT FORMAT_DATE('%d.%m.%Y', MAX(snapshot_date)) FROM `{P}.ozon_raw.RAW_OZON_STOCKS`) AS ozon_stock,
      (SELECT FORMAT_TIMESTAMP('%d.%m.%Y %H:%M', MAX(extracted_at), 'Europe/Moscow') FROM `{P}.ozon_raw.RAW_OZON_SUPPLY_ORDERS`) AS ozon_orders,
      (SELECT ANY_VALUE(plan_version) FROM `{P}.wb_mart.V_CT_PLAN_ACTIVE`) AS plan_version,
      (SELECT FORMAT_TIMESTAMP('%d.%m.%Y %H:%M', MAX(recorded_at), 'Europe/Moscow') FROM `{P}.evetis_ops.CT_STOCK_MOVEMENT`) AS ledger_last,
      (SELECT ANY_VALUE(config_value) FROM `{P}.evetis_ops.OPS_CONFIG` WHERE config_key = 'writeback_enabled') AS writeback''')[0]

    wb_cal = {c['channel']: c for c in cal}
    arr = {ch: next((r['arrival_date'] for r in plan if r['channel'] == ch), '') for ch in ('WB', 'OZON')}
    narr = {ch: next((r['next_arrival_date'] for r in plan if r['channel'] == ch), '') for ch in ('WB', 'OZON')}
    now = datetime.now(timezone(timedelta(hours=3))).strftime('%d.%m.%Y %H:%M')
    wbk = Workbook()

    # ---------- 01_SUPPLY_PLAN ----------
    ws = wbk.active
    ws.title = '01_SUPPLY_PLAN'
    s = Sheet(ws)
    for r in plan:
        r['kind'] = 'Набор' if r['is_bundle'] else 'Соло'
        r['reason_short'] = (r['demand_warning'] + ' · ' if r['demand_warning'] else '') + (r['reason'] or '')
    ship_rows = [r for r in plan if (r['rec_final'] or 0) > 0]
    hold_rows = [r for r in plan if (r['rec_final'] or 0) == 0 and r['status_code'] in ('REVIEW', 'FF LIMIT', 'WAIT')]
    wb_ship = [r for r in ship_rows if r['channel'] == 'WB']
    oz_ship = [r for r in ship_rows if r['channel'] == 'OZON']

    def pos(rows):
        return sum(r['rec_final'] or 0 for r in rows)

    def phys(rows):
        return sum(r['rec_physical_units'] or 0 for r in rows)

    assemble_pos = sum(b['to_assemble_now'] or 0 for b in bundles)
    assemble_res = sum(b['to_assemble_reserved'] or 0 for b in bundles)
    assemble_phys = sum((b['to_assemble_now'] or 0) * (b['components'] or 0) for b in bundles)
    move_new = sum(r['to_pick_from_pallet'] or 0 for r in pick)
    move_res = sum(r['reserved_on_pallet_to_pick'] or 0 for r in pick)

    s.title('EVETIS OPERATIONS · План поставок', [
        f'ТОЛЬКО ЧТЕНИЕ · данные evetis_ops на {now} МСК · отгрузка с ФФ {wb_cal["WB"]["ship_date"]} → '
        f'прибытие WB {arr["WB"]} / Ozon {arr["OZON"]} · следующая поставка {wb_cal["WB"]["next_ship_date"]} → '
        f'{narr["WB"]} / {narr["OZON"]} · покрытие 45 дней считается ПОСЛЕ прибытия, страховой отдельно.',
    ])
    s.band('ИТОГО СЕГОДНЯ', C_DARK, note='Детали — в блоках ниже; полный расчёт каждой строки — в конце листа.', note_as_comment=True)
    sum_cols = [col('Что сделать', 'what', 34), col('Позиции, шт.', 'positions', 13, INT),
                col('Физические\nединицы, шт.', 'units', 14, INT), col('Комментарий', 'comment', 110)]
    s.header(sum_cols)
    s.rows(sum_cols, [
        {'what': 'Отгрузить на WILDBERRIES', 'positions': pos(wb_ship), 'units': phys(wb_ship),
         'comment': f'{len(wb_ship)} позиций к отгрузке · прибытие {arr["WB"]}'},
        {'what': 'Отгрузить на OZON', 'positions': pos(oz_ship), 'units': phys(oz_ship),
         'comment': f'{len(oz_ship)} позиций к отгрузке · прибытие {arr["OZON"]}'},
        {'what': 'Снять с паллет на полку', 'positions': None, 'units': move_new + move_res,
         'comment': f'{move_new} под новые отгрузки + {move_res} под уже зарезервированные поставки Ozon от 07.09'},
        {'what': 'Собрать наборы', 'positions': assemble_pos, 'units': assemble_phys,
         'comment': f'в т.ч. {assemble_res} наборов под уже зарезервированные отгрузки'},
        {'what': 'Требуют вашего решения', 'positions': len(hold_rows), 'units': None,
         'comment': 'ПРОВЕРИТЬ / ФФ не хватает / ЖДАТЬ — блок «ПРОВЕРИТЬ / НЕ ОТГРУЖАТЬ» ниже'},
    ])

    s.band(f'ОТГРУЗИТЬ WB СЕЙЧАС   ·   {pos(wb_ship)} позиций / {phys(wb_ship)} физических единиц', C_WB,
           note='Лимит одной поставки WB через ПВЗ: 25 кг / 500 ед. / 200 л.', note_as_comment=True)
    s.header(TOP_COLS)
    s.rows(TOP_COLS, sorted(wb_ship, key=lambda r: -(r['rec_physical_units'] or 0)), status_key='status_label',
           bold_key='rec_final', total={'product_name': 'ИТОГО', 'rec_final': pos(wb_ship),
                                        'rec_physical_units': phys(wb_ship)})
    s.band(f'ОТГРУЗИТЬ OZON СЕЙЧАС   ·   {pos(oz_ship)} позиций / {phys(oz_ship)} физических единиц', C_OZ,
           note='ON_OZON (остаток площадки), IN_ACCEPTANCE (на приёмке) и RESERVED_OZON_ON_FF (резерв на ФФ) — '
                'разные вещи и не складываются.', note_as_comment=True)
    s.header(TOP_COLS)
    s.rows(TOP_COLS, sorted(oz_ship, key=lambda r: -(r['rec_physical_units'] or 0)), status_key='status_label',
           bold_key='rec_final', total={'product_name': 'ИТОГО', 'rec_final': pos(oz_ship),
                                        'rec_physical_units': phys(oz_ship)})
    s.band(f'ПРОВЕРИТЬ / НЕ ОТГРУЖАТЬ   ·   {len(hold_rows)} строк', 'A61C00',
           note='Количество не уменьшается автоматически: решение принимает владелец.', note_as_comment=True)
    s.header(TOP_COLS)
    s.rows(TOP_COLS, sorted(hold_rows, key=lambda r: (r['status_code'], r['channel'], r['card_sku'])),
           status_key='status_label')

    s.band('ТЗ ДЛЯ ФУЛФИЛМЕНТА (Usend) — можно копировать в задание складу', C_PICK,
           note='Считается из плана выше. В журнал ничего не пишется: это расчёт, а не операция.')
    ws.cell(s.r, 1, '1. Снять с паллет на полку').font = Font(FONT, 11, bold=True, color=C_PICK)
    s.r += 1
    s.header(TASK_MOVE_COLS)
    s.rows(TASK_MOVE_COLS, [r for r in task if (r['move_pallet_to_shelf_total'] or 0) > 0], total={
        'product_name': 'ИТОГО', 'move_pallet_to_shelf_new': move_new, 'move_pallet_to_shelf_reserved': move_res,
        'move_pallet_to_shelf_total': move_new + move_res})
    ws.cell(s.r, 1, '2. Собрать наборы').font = Font(FONT, 11, bold=True, color=C_PICK)
    s.r += 1
    s.header(TASK_BUILD_COLS)
    s.rows(TASK_BUILD_COLS, [dict(b, assemble_physical_units=(b['to_assemble_now'] or 0) * (b['components'] or 0))
                             for b in bundles if (b['to_assemble_now'] or 0) > 0],
           total={'product_name': 'ИТОГО', 'to_assemble_now': assemble_pos, 'to_assemble_reserved': assemble_res,
                  'assemble_physical_units': assemble_phys})
    ws.cell(s.r, 1, '3. Расход компонентов и что остаётся на ФФ').font = Font(FONT, 11, bold=True, color=C_PICK)
    s.r += 1
    s.header(TASK_COMP_COLS)
    s.rows(TASK_COMP_COLS, [r for r in task if (r['total_physical_units'] or 0) > 0 or (r['reserved_units_already'] or 0) > 0],
           total={'product_name': 'ИТОГО', **{k: sum(r[k] or 0 for r in task) for k in
                  ('solo_units_wb', 'solo_units_ozon', 'bundle_component_units', 'fbs_component_units',
                   'total_physical_units', 'reserved_units_already', 'free_ff_after_operation')}})

    s.band('ПОДРОБНЫЙ РАСЧЁТ (для проверки)', C_DARK,
           note='Полная таблица по каждой строке: позиция запаса, спрос до прибытия, цель, потребность, округление, ворота.')
    for channel, color, name in (('WB', C_WB, 'WILDBERRIES'), ('OZON', C_OZ, 'OZON')):
        c = wb_cal[channel]
        rows = [r for r in plan if r['channel'] == channel]
        s.band(f'{name}   ·   отгрузка {c["ship_date"]} → прибытие {arr[channel]}   ·   '
               f'покрытие {c["target_cover_days"]} дн + страховой {c["safety_stock_days"]} дн   ·   '
               f'срок канала {c["lead_time_days"]} дн   ·   к отгрузке {pos(rows)} позиций / {phys(rows)} физ. ед.', color)
        s.header(PLAN_COLS)
        s.rows(PLAN_COLS, rows, status_key='status_label', bold_key='rec_final',
               total={'product_name': 'ИТОГО', 'rec_final': pos(rows), 'rec_physical_units': phys(rows),
                      'on_marketplace': sum(r['on_marketplace'] or 0 for r in rows),
                      'committed_inbound': sum(r['committed_inbound'] or 0 for r in rows),
                      'need_math': sum(r['need_math'] or 0 for r in rows)})
    s.band('BUNDLE PRODUCTION · сборка наборов', C_BUNDLE, note=(
        'Мощности наборов не складываются: наборы делят компоненты. «Собрать сейчас» = под уже зарезервированные '
        'отгрузки + новые рекомендации + резерв FBS − уже собранное.'))
    s.header(BUNDLE_COLS)
    s.rows(BUNDLE_COLS, bundles, bold_key='to_assemble_now',
           total={'product_name': 'ИТОГО', 'to_assemble_now': assemble_pos,
                  'ship_wb': sum(b['ship_wb'] or 0 for b in bundles),
                  'ship_ozon': sum(b['ship_ozon'] or 0 for b in bundles), 'to_assemble_reserved': assemble_res})
    s.band('СОБРАТЬ СЕЙЧАС → компоненты (раскладка по BOM)', C_BUNDLE)
    s.header(BOM_COLS)
    s.rows(BOM_COLS, bom, total={'product_name': 'ИТОГО',
                                 'units_total': sum(b['units_total'] or 0 for b in bom),
                                 'units_from_reserved': sum(b['units_from_reserved'] or 0 for b in bom),
                                 'units_from_free': sum(b['units_from_free'] or 0 for b in bom)})
    s.band('PICK FROM STORAGE · снять с паллет', C_PICK, note=(
        'Итог спроса — ровно сумма рекомендаций выше: соло WB + соло Ozon + компоненты наборов + FBS. '
        'Полка после снятия не может быть отрицательной.'))
    s.header(PICK_COLS)
    s.rows(PICK_COLS, pick, total={
        'product_name': 'ИТОГО',
        'solo_wb_need': sum(r['solo_wb_need'] or 0 for r in pick),
        'solo_ozon_need': sum(r['solo_ozon_need'] or 0 for r in pick),
        'bundle_component_need': sum(r['bundle_component_need'] or 0 for r in pick),
        'total_physical_demand': sum(r['total_physical_demand'] or 0 for r in pick),
        'to_pick_from_pallet': move_new, 'reserved_on_pallet_to_pick': move_res,
        'pick_total_with_reserved': sum(r['pick_total_with_reserved'] or 0 for r in pick)})
    ws.freeze_panes = 'C6'

    # ---------- 02_SHIPMENTS ----------
    ws = wbk.create_sheet('02_SHIPMENTS')
    s = Sheet(ws)
    s.title('Отгрузки и поставки', [
        f'Данные на {now} (МСК). Один заказ — один этап: резерв на ФФ, приёмка, в пути или принято. '
        'Единица не может быть одновременно в двух этапах.'], width=17)
    s.band('OZON · поставки', C_OZ, width=17)
    s.header(SHIP_COLS)
    s.rows(SHIP_COLS, ships, total={'channel': 'ИТОГО', 'cards': sum(r['cards'] or 0 for r in ships),
                                    'physical_units': sum(r['physical_units'] or 0 for r in ships),
                                    'reserved_on_ff_units': sum(r['reserved_on_ff_units'] or 0 for r in ships),
                                    'pipeline_units': sum(r['pipeline_units'] or 0 for r in ships),
                                    'handed_over_units': sum(r['handed_over_units'] or 0 for r in ships)})
    ws.cell(s.r, 1, 'Поставок WB в работе нет: последняя поставка 22.07.2026, черновики в кабинете поставками не считаются.')
    ws.cell(s.r, 1).font = Font(FONT, 9, italic=True, color='666666')
    ws.freeze_panes = 'C6'

    # ---------- 03_FF_STOCK ----------
    ws = wbk.create_sheet('03_FF_STOCK')
    s = Sheet(ws)
    total_ff = sum(r['total_physical'] or 0 for r in ff)
    s.title('Запас на фулфилменте (физические единицы)', [
        f'Источник правды — журнал evetis_ops на {now} (МСК). Старое число Control Tower как физическую правду не использовать.',
        f'ИТОГО НА ФФ: {total_ff:,} единиц'.replace(',', ' '),
        'SHIPPED UNCONFIRMED — уже уехало с ФФ и в итог не входит. ON_OZON / ON_WB — остатки площадок, показаны справочно.',
    ], width=14)
    s.band('СОСТОЯНИЕ ПО SKU', C_PICK, width=14)
    s.header(FF_COLS)
    s.rows(FF_COLS, ff, total={'product_name': 'ИТОГО', **{
        k: sum(r[k] or 0 for r in ff) for k in
        ('pallet_free', 'shelf_free', 'reserved_wb', 'reserved_ozon', 'reserved_ozon_on_pallet', 'assembled',
         'fbs_ready', 'shipped_unconfirmed', 'total_physical', 'ozon_in_acceptance', 'on_ozon', 'on_wb')}})
    ws.freeze_panes = 'C8'

    # ---------- 04_SETTINGS ----------
    ws = wbk.create_sheet('04_SETTINGS')
    s = Sheet(ws)
    s.title('Настройки и правила (только чтение)', [
        'Значения живут в BigQuery (evetis_ops.REF_CHANNEL_SHIPPING, REF_SKU_LOGISTICS, OPS_CONFIG). '
        'Менять их в листе бессмысленно: при обновлении лист берёт значения из BigQuery.'], width=13)
    s.band('ПРАВИЛА РАСЧЁТА (одобрены владельцем)', C_DARK, width=13)
    for line in [
        '1. Целевое покрытие 45 дней — ПОСЛЕ прибытия. Сборка ФФ (2 рабочих дня) и срок канала учитываются отдельно и видны в листе.',
        '2. Округление вверх до операционной кратности. Потребность меньше кратности → ЖДАТЬ, если страховой запас доживёт '
        'до следующей возможной поставки (недельный цикл); иначе отгружается один минимум.',
        '3. Ozon при низком спросе: одиночные 5, наборы 2 — только если потребность меньше обычной кратности и не меньше уменьшенной.',
        '4. ПРОВЕРИТЬ, если после округления покрытие выше: WB 66 дней (45 + 7 + 14), Ozon 69 дней (45 + 10 + 14); '
        'или если запас не распродаётся минимум за 30 дней до срока годности. Количество при этом не уменьшается — решает владелец.',
        '5. Позиция запаса: остаток площадки + подтверждённое входящее + приёмка. Каждая единица считается один раз.',
        '6. Ozon: ON_OZON, IN_ACCEPTANCE и RESERVED_OZON_ON_FF — раздельно. Приёмка никогда не прибавляется к остатку площадки. '
        'Если неизвестное принятое количество может изменить решение «отгрузить ↔ ждать» — статус ПРОВЕРИТЬ.',
        '7. ФФ — один общий пул: WB и Ozon не могут претендовать на одну и ту же единицу.',
    ]:
        c = ws.cell(s.r, 1, line)
        c.font = Font(FONT, 10, color=C_DARK)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.merge_cells(start_row=s.r, start_column=1, end_row=s.r, end_column=13)
        ws.row_dimensions[s.r].height = 28
        s.r += 1
    s.r += 1
    s.band('КАНАЛ И СПОСОБ ОТГРУЗКИ', C_DARK, width=13)
    ch_cols = [col('Канал', 'channel', 10), col('Способ', 'shipping_method', 22),
               col('Покрытие,\nдн', 'target_cover_days', 10, INT), col('Страховой,\nдн', 'safety_stock_days', 10, INT),
               col('Срок,\nдн', 'lead_time_days', 9, INT), col('ФФ, раб. дн', 'ff_prep_working_days', 10, INT),
               col('Лимит,\nкг', 'max_lot_weight_kg', 9, INT), col('Лимит,\nед.', 'max_lot_units', 9, INT),
               col('Лимит,\nл', 'max_lot_volume_l', 9, INT), col('Действует с', 'effective_from', 12, DATE),
               col('Источник', 'source', 30)]
    s.header(ch_cols)
    s.rows(ch_cols, chan)
    s.band('ПАРАМЕТРЫ C1', C_DARK, width=13)
    cfg_cols = [col('Ключ', 'config_key', 30), col('Значение', 'config_value', 12), col('Что это', 'note', 110)]
    s.header(cfg_cols)
    s.rows(cfg_cols, cfg)
    s.band('ЛОГИСТИКА ПО SKU × КАНАЛ', C_DARK, width=13, note=(
        'Заводской короб не заполняется, пока ФФ не подтвердит: показана только закономерность из данных как доказательство.'))
    lg_cols = [col('SKU', 'internal_sku', 30), col('Канал', 'channel', 8),
               col('Кратность', 'shipment_multiple', 10, INT), col('Низкий спрос', 'shipment_multiple_low_demand', 11, INT),
               col('Минимум', 'min_shipment_units', 9, INT), col('Заводской короб', 'factory_carton_qty', 13, INT),
               col('Вес, кг', 'unit_weight_kg', 9, F1), col('Объём, л', 'unit_volume_l', 9, F1),
               col('Резерв FBS', 'fbs_reserve_units', 9, INT),
               col('Вес: класс', 'weight_class', 11), col('Кратность: класс', 'multiple_class', 15),
               col('Короб: доказательство', 'carton_hint', 60), col('Действует с', 'effective_from', 12, DATE)]
    s.header(lg_cols)
    s.rows(lg_cols, logi)
    s.band('СВЕЖЕСТЬ ДАННЫХ', C_DARK, width=13)
    for k, v in (('Остатки WB (снимок)', fresh['wb_stock']), ('Остатки Ozon (снимок)', fresh['ozon_stock']),
                 ('Заказы поставок Ozon (выгрузка)', fresh['ozon_orders']), ('Версия плана продаж', fresh['plan_version']),
                 ('Последнее движение в журнале ФФ', fresh['ledger_last']),
                 ('Запись владельца (writeback_enabled)', fresh['writeback'])):
        ws.cell(s.r, 1, k).font = Font(FONT, 10, color=C_DARK)
        ws.cell(s.r, 3, v).font = Font(FONT, 10, bold=True, color=C_DARK)
        s.r += 1

    # ---------- DATA (скрытый) ----------
    ws = wbk.create_sheet('DATA')
    r = 1
    ws.cell(r, 1, 'Сырые строки представлений evetis_ops — источник чисел этой книги. Лист скрыт.').font = Font(FONT, 10, bold=True)
    r += 2
    for name, rows in (('V_OPS_SUPPLY_PLAN', plan), ('V_OPS_BUNDLE_PRODUCTION', bundles),
                       ('V_OPS_BUNDLE_BOM_EXPANSION', bom), ('V_OPS_PICK_FROM_STORAGE', pick),
                       ('V_OPS_FF_STOCK_SHEET', ff), ('V_OPS_SHIPMENTS_SHEET', ships), ('V_OPS_FF_TASK', task),
                       ('V_OPS_SUPPLY_CALENDAR', cal)):
        ws.cell(r, 1, name).font = Font(FONT, 10, bold=True, color=C_WB)
        r += 1
        if rows:
            keys = list(rows[0].keys())
            for i, k in enumerate(keys, start=1):
                c = ws.cell(r, i, k)
                c.font = Font(FONT, 9, bold=True, color=C_WHITE)
                c.fill = PatternFill('solid', fgColor=C_DARK)
            r += 1
            for rec in rows:
                for i, k in enumerate(keys, start=1):
                    v = rec[k]
                    ws.cell(r, i, v if not isinstance(v, bool) else str(v))
                r += 1
        r += 2
    ws.sheet_state = 'hidden'

    wbk.save(path)
    return path, {'plan_rows': len(plan), 'ship_total': sum(r['rec_final'] or 0 for r in plan),
                  'bundles': len(bundles), 'pick': len(pick), 'ff_total': total_ff, 'shipments': len(ships)}


if __name__ == '__main__':
    out = sys.argv[1] if len(sys.argv) > 1 else 'EVETIS_OPERATIONS.xlsx'
    path, stats = build(out)
    print('saved', path, os.path.getsize(path), 'bytes', stats)
