#!/usr/bin/env python3
"""EVETIS OPERATIONS · C2 — фикстуры для tools/ops_c2_logic_test.js.

Сохраняет в <каталог>:
  <V_OPS_SHEET_*>.json  — сырые ответы BigQuery (schema + rows в формате jobs.getQueryResults), как их видит Apps Script;
  live_colA.json        — тексты колонки A листов 01–05 из экспорта книги (.xlsx), для проверки установки раскладки.

Фикстуры содержат рабочие данные и в Git не кладутся.

Запуск:  python3 tools/ops_c2_fixtures.py <каталог> [<экспорт книги .xlsx>]
Экспорт книги: Drive → EVETIS OPERATIONS → Скачать → .xlsx (или Drive API files.export).
"""
import json
import os
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import ops_deploy as d  # noqa: E402

VIEWS = ['V_OPS_SHEET_PLAN', 'V_OPS_SHEET_BUNDLES', 'V_OPS_SHEET_BOM', 'V_OPS_SHEET_PICK', 'V_OPS_SHEET_FF_TASK',
         'V_OPS_SHEET_FF_STOCK', 'V_OPS_SHEET_SHIPMENTS', 'V_OPS_SHEET_CHANNELS', 'V_OPS_SHEET_CONFIG',
         'V_OPS_SHEET_LOGISTICS', 'V_OPS_SHEET_META']
SHEETS = ['01_SUPPLY_PLAN', '02_SHIPMENTS', '03_FF_STOCK', '04_SETTINGS', '05_РАСЧЁТ']


def raw_query(sql):
    r = d.call('POST', f'{d.API}/jobs', {
        'configuration': {'query': {'query': sql, 'useLegacySql': False},
                          'labels': {'stage': 'stage-c2', 'sandbox': 'false'}},
        'jobReference': {'projectId': d.P, 'location': 'EU'}})
    if 'error' in r:
        raise SystemExit('BQ insert: ' + str(r['error'])[:300])
    jid = r['jobReference']['jobId']
    while True:
        j = d.call('GET', f'{d.API}/jobs/{jid}?location=EU')
        if j.get('status', {}).get('state') == 'DONE':
            break
        time.sleep(1.5)
    if j['status'].get('errorResult'):
        raise SystemExit('BQ query: ' + j['status']['errorResult'].get('message', '')[:400])
    rows, token, schema = [], None, None
    while True:
        res = d.call('GET', f'{d.API}/queries/{jid}?location=EU&maxResults=5000' + (f'&pageToken={token}' if token else ''))
        schema = res.get('schema')
        rows += res.get('rows', [])
        token = res.get('pageToken')
        if not token:
            return {'schema': schema, 'rows': rows, 'totalRows': str(len(rows)), 'jobComplete': True}


def main(out, xlsx=None):
    os.makedirs(out, exist_ok=True)
    for v in VIEWS:
        res = raw_query(f'SELECT * FROM `{d.P}.evetis_ops.{v}`')
        json.dump(res, open(os.path.join(out, v + '.json'), 'w'), ensure_ascii=False)
        print(f'{v}: {len(res["rows"])} строк')
    if xlsx:
        import openpyxl
        wb = openpyxl.load_workbook(xlsx)
        col_a = {s: ['' if wb[s].cell(r, 1).value is None else str(wb[s].cell(r, 1).value)
                     for r in range(1, wb[s].max_row + 1)] for s in SHEETS}
        json.dump(col_a, open(os.path.join(out, 'live_colA.json'), 'w'), ensure_ascii=False)
        print('live_colA.json:', {k: len(v) for k, v in col_a.items()})
    return 0


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    sys.exit(main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else None))
