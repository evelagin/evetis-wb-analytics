import glob, os, re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

UP = "/root/.claude/uploads/f5f0e910-5cf1-5396-8501-ec7b791bef4a"
OUT = "/mnt/user-data/outputs/tz_ozon"
os.makedirs(OUT, exist_ok=True)

# ---- справочник (ff-tz-fill-rules + REF_PRODUCT_MASTER) ----
SINGLE = {  # nm_id: (название в ТЗ, баркод, вариант требований)
 "252442517": ("крем для рук", "2040927586377", "V"),
 "593111986": ("крем для рук и тела Вишня", "2047045125709", "B"),
 "593111985": ("крем для рук и тела Амбра", "2047045119074", "B"),
 "305101272": ("сыворотка для лица увлажняющая", "2042315776735", "V"),
 "305101361": ("сыворотка для лица анти-акне", "2042315779385", "V"),
 "438775437": ("крем для лица увлажняющий", "2044356263243", "V"),
 "438775617": ("крем для лица анти-акне", "2044356266398", "V"),
 "535581674": ("тоник для лица увлажняющий", "2046172237057", "V"),
 "535581675": ("тоник для лица анти-акне", "2046172237064", "V"),
 "535580776": ("пудра энзимная", "2046172224798", "V"),
}
BUNDLE = {  # nm_id набора: (название, баркод набора, [nm_id компонентов])
 "930334395": ("набор: крем для рук + крем Вишня", "2050118820111", ["252442517","593111986"]),
 "930334396": ("набор: крем для рук + крем Амбра", "2050119291156", ["252442517","593111985"]),
 "930334397": ("набор: крем Вишня + крем Амбра", "2050119289283", ["593111986","593111985"]),
 "910330849": ("набор увл: тоник + сыворотка", "2049908507615", ["535581674","305101272"]),
 "1083392113":("набор акне: тоник + сыворотка", "2051601151057", ["535581675","305101361"]),
 "952068582": ("набор акне: тоник + крем", "2050315992109", ["535581675","438775617"]),
 "567668635": ("набор акне: сыворотка + крем", "2046706658488", ["305101361","438775617"]),
 "773170316": ("набор увл: тоник + сыворотка + крем", "2048495880361", ["535581674","305101272","438775437"]),
 "910584041": ("набор акне: тоник + сыворотка + крем", "2049910829446", ["535581675","305101361","438775617"]),
 "868597351": ("набор акне: пудра + тоник + сыворотка + крем", "2049477631605", ["535580776","535581675","305101361","438775617"]),
 "909951444": ("набор акне: пудра + сыворотка + крем", "2049904839437", ["535580776","305101361","438775617"]),
}

def text_bundle(nm, comps):
    name, bc, cl = BUNDLE[nm]
    lines = [f"{i+1}) Взять «{SINGLE[c][0]}», баркод {SINGLE[c][1]}." for i, c in enumerate(cl)]
    k = len(cl)
    lines += [f"{k+1}) Соединить позиции скотчем — собрать набор.",
              f"{k+2}) Наклеить ШК набора: баркод {bc}, артикул {nm}.",
              f"{k+3}) Наклеить КИЗ набора из файла, который мы пришлём.",
              f"{k+4}) Прислать нам Excel-файл сборки: столбец A — КИЗ набора, столбцы B, C, D… — КИЗ каждого компонента, вошедшего именно в этот набор. Одна строка = один собранный набор."]
    return "\n".join(lines)

TEXT_B = ("1) Достать крем из коробки.\n"
          "2) Наклеить КИЗ из файла, который мы пришлём, закрыв КИЗ, уже нанесённый на упаковку.\n"
          "3) Прислать нам Excel-файл: столбец A — КИЗ, который наклеили на единицу. Одна строка = одна единица.")
TEXT_V = "Товар промаркирован и готов к отгрузке. Дополнительных работ не требуется."

# ---- стили ----
F = "Arial"
thin = Side(style="thin", color="000000")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
HFILL = PatternFill("solid", fgColor="D9E1F2")
LFILL = PatternFill("solid", fgColor="F2F2F2")

def put(ws, ref, v, bold=False, wrap=True, align="left", fill=None, border=True, size=10):
    c = ws[ref]; c.value = v
    c.font = Font(name=F, size=size, bold=bold)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if fill: c.fill = fill
    if border: c.border = box
    return c

def build(path):
    wb_src = load_workbook(path, data_only=True)
    ws_src = wb_src.active
    rows = [r for r in ws_src.iter_rows(min_row=2, values_only=True) if r[0]]
    supply = str(rows[0][0]); cluster = str(rows[0][1])
    items = [(str(r[3]), int(r[5])) for r in rows]   # (артикул, кол-во)
    # порядок: сначала одиночные, потом наборы
    items.sort(key=lambda x: (x[0] in BUNDLE, x[0]))

    wb = Workbook(); ws = wb.active; ws.title = "Техническое задание"
    # реквизиты
    for r in range(1, 13):
        ws.merge_cells(f"B{r}:C{r}"); ws.merge_cells(f"D{r}:F{r}")
    head = {1:("Заказчик","ИП Елагина"), 2:("Юридическое лицо","ИП Елагина"),
            3:("Номер поставки маркетплейса", supply), 4:("Маркетплейс","озон"),
            5:("Склад / кластер назначения", f"Ozon · {cluster}"), 6:("Дата отгрузки", "04.09.2026"),
            7:("Адрес сдачи",""), 8:("Тип упаковки","короб"), 9:("Количество коробов", 1),
            10:("Количество паллет",""), 11:("Количество товара в поставке (отгрузочных единиц)", sum(q for _,q in items)),
            12:("Номер заказа USEND","")}
    for r,(k,v) in head.items():
        put(ws, f"B{r}", k, bold=True, fill=LFILL); put(ws, f"D{r}", v)
    ws.merge_cells("K4:L4"); ws.merge_cells("K5:L5")
    put(ws, "K4", "Видеофиксация", bold=True, fill=LFILL); put(ws, "M4", "НЕТ")
    put(ws, "K5", "Работа с КИЗ", bold=True, fill=LFILL); put(ws, "M5", "ДА")
    put(ws, "H1", "ТЕХНИЧЕСКОЕ ЗАДАНИЕ НА ПОСТАВКУ OZON FBO", bold=True, border=False, size=13, wrap=False)

    HDR = ["№ п/п","Usend ID товара","Название товара","Зона размещения","Размер","Цвет",
           "Артикул","Количество (план)","Баркод","Количество (факт)","Номер короба","Номер паллета",
           "Наличие ШК","Требования к упаковке и доп. работам","Требуется маркировка КИЗ",
           "Номер КИЗ","Ссылка на файл КИЗ"]
    for i,h in enumerate(HDR):
        put(ws, f"{chr(65+i)}14", h, bold=True, align="center", fill=HFILL)
    ws.row_dimensions[14].height = 45

    r = 15; n = 0; ship = {}
    for nm, qty in items:
        n += 1
        if nm in SINGLE:
            name, bc, var = SINGLE[nm]
            vals = {"A":n,"C":name,"G":nm,"H":qty,"I":bc,"K":1,"M":"есть шк",
                    "N": TEXT_B if var=="B" else TEXT_V, "O": "да" if var=="B" else "нет"}
            for col in "ABCDEFGHIJKLMNOPQ":
                put(ws, f"{col}{r}", vals.get(col), align="center" if col not in "CN" else "left")
            ws.row_dimensions[r].height = 60 if var=="B" else 20
            ship[bc] = ship.get(bc,0)+qty
            r += 1
        else:
            name, bc, comps = BUNDLE[nm]
            top = r
            for j, c in enumerate(comps):
                vals = {"A":n if j==0 else n, "C":SINGLE[c][0], "G":c, "H":qty}
                if j == 0: vals.update({"J":None,"K":1,"O":"да"})
                for col in "ABCDEFGHIJKLMNOPQ":
                    put(ws, f"{col}{r}", vals.get(col), align="center" if col not in "CN" else "left")
                r += 1
            # № п/п сквозной по каждой строке компонента
            for j in range(len(comps)):
                if j > 0:
                    n += 1; ws[f"A{top+j}"].value = n
            bottom = r-1
            for col, v in (("I", bc), ("M", "наклеить шк"), ("N", text_bundle(nm, comps))):
                ws.merge_cells(f"{col}{top}:{col}{bottom}")
                put(ws, f"{col}{top}", v, align="center" if col!="N" else "left")
            for rr in range(top, bottom+1):
                ws.row_dimensions[rr].height = max(30, 110 // len(comps))
            ship[bc] = ship.get(bc,0)+qty

    widths = dict(A=6,B=10,C=30,D=8,E=7,F=7,G=12,H=10,I=15,J=10,K=9,L=9,M=12,N=70,O=12,P=10,Q=12)
    for k,v in widths.items(): ws.column_dimensions[k].width = v
    ws.freeze_panes = "A15"
    ws.page_setup.orientation = "landscape"; ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0

    # ---- Отгрузочный лист ----
    ws2 = wb.create_sheet("Отгрузочный лист")
    for r2,(k,v) in enumerate([("Номер заказа USEND",""),("Тип отгрузки","короб"),("Паллет",""),
                               ("Коробов",1),("Итого, шт", sum(ship.values()))], start=1):
        put(ws2, f"A{r2}", k, bold=True, fill=LFILL); put(ws2, f"B{r2}", v)
    put(ws2, "A6", f"Поставка Ozon {supply} · {cluster}", bold=True, border=False)
    for i,h in enumerate(["Название","Баркод","Количество (факт)","Номер короба"]):
        put(ws2, f"{chr(65+i)}7", h, bold=True, align="center", fill=HFILL)
    names = {bc: SINGLE[k][0] for k in SINGLE for bc in [SINGLE[k][1]]}
    names.update({BUNDLE[k][1]: BUNDLE[k][0] for k in BUNDLE})
    r2 = 8
    for bc, q in ship.items():
        put(ws2, f"A{r2}", names[bc]); put(ws2, f"B{r2}", bc, align="center")
        put(ws2, f"C{r2}", q, align="center"); put(ws2, f"D{r2}", 1, align="center"); r2 += 1
    put(ws2, f"A{r2}", "ИТОГО", bold=True); put(ws2, f"C{r2}", f"=SUM(C8:C{r2-1})", bold=True, align="center")
    for k,v in dict(A=42,B=16,C=18,D=14).items(): ws2.column_dimensions[k].width = v

    safe = re.sub(r"[^\w]+", "_", cluster).strip("_")
    out = f"{OUT}/ТЗ_USEND_Ozon_{supply}_{safe}.xlsx"
    wb.save(out)
    return out, supply, cluster, sum(q for _,q in items), len(items), sum(q for nm,q in items if nm in BUNDLE)

for p in sorted(glob.glob(f"{UP}/*2000065*.xlsx")):
    print(build(p))
